"""Unit tests for publish.py (synthetic results, in-process fake Nextcloud).
python3 -m unittest discover -s terraform/lxc/ansible/files/eval-runner -p "test_*.py"
"""

import io
import json
import os
import tempfile
import unittest

import mock_nextcloud
import publish

try:
    import openpyxl
except ImportError:  # the image has it; locally: pip install openpyxl==3.1.5 in a venv
    openpyxl = None
import summarize
from test_summarize import gpqa_rows, ifeval_rows, write_run


class FakeNextcloud:
    """publish.Nextcloud's interface, backed by mock_nextcloud's state and routing."""

    def __init__(self):
        mock_nextcloud.reset()
        self.state = mock_nextcloud.STATE
        self.calls = []

    def ensure_folder(self, rel):
        parts = rel.strip("/").split("/")
        for i in range(1, len(parts) + 1):
            path = "/".join(parts[:i])
            if path not in self.state["folders"]:
                self.state["folders"].append(path)

    def request(self, method, path, body=None):
        self.calls.append((method, path))
        assert path.startswith(mock_nextcloud.OCS_API), path
        code, result = mock_nextcloud.ocs_route(method, path[len(mock_nextcloud.OCS_API):], body or {})
        if code != 200:
            raise publish.NextcloudError(f"{method} {path} -> {code} {result}")
        return result

    def put_file(self, rel, content):
        assert os.path.dirname(rel) in self.state["folders"], rel
        self.state["files"][rel] = mock_nextcloud.stored_content(content)

    def delete_file(self, rel):
        self.state["deleted"].append(rel)
        self.state["files"].pop(rel, None)

    def tables(self, method, path, body=None):
        self.calls.append((method, path))
        code, result = mock_nextcloud.route(method, path.split("?", 1)[0], body or {})
        if code != 200:
            raise publish.NextcloudError(f"{method} {path} -> {code} {result}")
        return result


def make_tree(root):
    """One eval-runner run (with run.json) and two historical sources (one excluded)."""
    run_dir = write_run(root, "glm-5.3-flash-both-20261002T000000Z", ["gpqa", "ifeval"],
                        gpqa_rows=gpqa_rows(["(A)", "", "(B)"]), ifeval_rows=ifeval_rows(["ok", ""]),
                        run_json=False)
    with open(os.path.join(run_dir, "run.json"), "w") as fh:
        json.dump({"run": os.path.basename(run_dir), "note": "reasoning_effort=high", "created_utc": "20261002T000000Z",
                   "fingerprint": "abcd" * 16, "limit": None, "concurrency": 1, "lm_eval_version": "0.4.12",
                   "server": {"model_id": "glm-5.3-flash", "props": {
                       "model_path": "/m/GLM-5.3-Flash-UD-IQ2_XXS-00001-of-00004.gguf",
                       "build_info": "b11309-a4d880fd5", "n_ctx": 131072,
                       "params": {"temperature": 1.0, "top_p": 0.95}}},
                   "run_metrics": {"source": "llama-server /metrics", "segments": 1, "duration_seconds": 5430.0,
                                   "model_under_test": {"prompt_tokens": 51234, "completion_tokens": 123456,
                                                        "prompt_seconds": 40.0, "generation_seconds": 6100.0,
                                                        "generation_tokens_per_second": 20.2,
                                                        "prompt_tokens_per_second": 1280.9}}}, fh)
    hist = os.path.join(root, summarize.HISTORICAL_DIR)
    os.makedirs(hist)
    write_run(hist, "qwen36-35b-redo", ["gpqa", "ifeval"], gpqa_rows=gpqa_rows(["(A)", ""]),
              ifeval_rows=ifeval_rows(["ok"]), run_json=False)
    write_run(hist, "qwen36-35b", ["ifeval"], ifeval_rows=ifeval_rows(["x"]), run_json=False,
              config={"limit": None, "gen_kwargs": {}, "model_args": {"model": "eval-qwen36", "base_url": "http://f:11434/v1"}})
    return run_dir


def add_32k_run(root):
    """A full GPQA run at the 32k budget: its own series, not comparable with 8k."""
    write_run(root, "glm-5.3-flash-gpqa-32k-20261003T000000Z", ["gpqa"], gpqa_rows=gpqa_rows(["(A)", "(B)"]),
              run_json=False, config={"limit": None, "gen_kwargs": {"max_gen_toks": 32768},
                                      "model_args": {"model": "glm-5.3-flash"}})


class CollectTest(unittest.TestCase):
    def test_rows_from_runs_and_history(self):
        with tempfile.TemporaryDirectory() as root:
            make_tree(root)
            rows = [r for _, _, _, rs in publish.collect(root) for r in rs]
        by_key = {r["Key"]: r for r in rows}
        self.assertEqual(len(rows), 5)
        glm = by_key["runs/glm-5.3-flash-both-20261002T000000Z/gpqa_diamond_cot_zeroshot"]
        self.assertEqual(glm["Model"], "glm-5.3-flash")
        self.assertEqual(glm["Runtime"], "llama.cpp b11309-a4d880fd5")
        self.assertEqual(glm["Model file / tag"], "GLM-5.3-Flash-UD-IQ2_XXS-00001-of-00004.gguf")
        self.assertEqual(glm["Note"], "reasoning_effort=high")
        self.assertEqual(glm["Date"], "2026-10-02")
        self.assertEqual((glm["Score %"], glm["Empty answers"], glm["Unparsed"]), (50.0, 1, 1))
        self.assertEqual(glm["Empty %"], 33.3)
        self.assertEqual(glm["Comparable"], "yes")
        self.assertEqual((glm["Duration (min)"], glm["Tokens generated"], glm["Tokens/s"]), (90.5, 123456, 20.2))
        bug6 = [r for r in rows if r["Run"] == "qwen36-35b"][0]
        self.assertEqual(bug6["Comparable"], "no")
        self.assertIn("Bug 6", bug6["Why not comparable"])
        self.assertEqual(bug6["Runtime"], "Ollama")
        self.assertEqual((bug6["Duration (min)"], bug6["Tokens/s"]), (None, None))
        self.assertTrue(bug6["Key"].startswith("historical/qwen36-35b/ifeval/"))


class RenderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        make_tree(self.tmp.name)
        self.files, self.rows = publish.build_files(publish.collect(self.tmp.name), "# findings\n", "NOW")

    def tearDown(self):
        self.tmp.cleanup()

    def test_expected_files(self):
        self.assertIn("leaderboard.md", self.files)
        self.assertNotIn("leaderboard.csv", self.files)
        if openpyxl:
            self.assertIn("leaderboard.xlsx", self.files)
        self.assertEqual(self.files["findings.md"], b"# findings\n")
        self.assertIn("runs/glm-5.3-flash-both-20261002T000000Z/report.md", self.files)
        self.assertIn("historical/qwen36-35b/manifest.json", self.files)
        self.assertFalse(any("samples_" in name for name in self.files))

    def test_markdown_tables_have_consistent_columns(self):
        for name, content in self.files.items():
            if not name.endswith(".md"):
                continue
            header_cols = None
            for line in content.decode().splitlines():
                if not line.startswith("|"):
                    header_cols = None
                    continue
                cols = line.count("|")
                header_cols = header_cols or cols
                self.assertEqual(cols, header_cols, f"{name}: {line}")

    def test_leaderboard_ranks_and_lists_excluded(self):
        board = self.files["leaderboard.md"].decode()
        self.assertIn("| 1 | glm-5.3-flash | 50.00% | 1/3 |", board)
        self.assertIn("## Not comparable", board)
        self.assertIn("no max_gen_toks=8192", board)

    def test_manifest_follows_convention(self):
        manifest = json.loads(self.files["runs/glm-5.3-flash-both-20261002T000000Z/manifest.json"])
        for key in ("project", "run_id", "started_at", "finished_at", "summary"):
            self.assertIn(key, manifest)
        self.assertEqual(manifest["project"], "eval-runner")

    def test_markdown_tables_stay_narrow(self):
        """No markdown table wider than 4 columns (Nextcloud scrolls wider ones)."""
        for name, content in self.files.items():
            if name.endswith(".md"):
                for line in content.decode().splitlines():
                    if line.startswith("|"):
                        self.assertLessEqual(line.count("|") - 1, 4, f"{name}: {line}")

    def test_report_has_metric_value_table_per_task(self):
        report = self.files["runs/glm-5.3-flash-both-20261002T000000Z/report.md"].decode()
        self.assertIn("### GPQA diamond\n\n| Metric | Value |", report)
        self.assertIn("| Score (flexible-extract) | 50.00% |", report)
        self.assertIn("| Strict-match | 25.00% |", report)
        self.assertIn("| Empty answers | 1 (33.3%) |", report)

    def test_report_has_run_metrics(self):
        report = self.files["runs/glm-5.3-flash-both-20261002T000000Z/report.md"].decode()
        self.assertIn("## Run metrics\n\n- **Duration:** 1h 30m 30s\n", report)
        self.assertIn("| Generated tokens (incl. reasoning) | 123,456 |", report)
        self.assertIn("| Generation speed (tokens/s) | 20.2 |", report)
        self.assertNotIn("## Run metrics", self.files["historical/qwen36-35b/report.md"].decode())

    def test_run_metrics_unavailable_or_missing(self):
        self.assertIn("unavailable: reloaded", "\n".join(publish.render_metrics(
            {"run_metrics": {"duration_seconds": 60, "unavailable": "reloaded"}})))
        self.assertIn("Not recorded", publish.render_metrics({})[0])

    def test_report_header_is_tidy(self):
        report = self.files["runs/glm-5.3-flash-both-20261002T000000Z/report.md"].decode()
        self.assertIn("- **Started:** 2026-10-02 00:00 UTC", report)
        self.assertIn("temperature 1, top_p 0.95", report)
        self.assertIn("- **Harness:** lm_eval 0.4.12; sample limit none (full run)", report)
        self.assertEqual(publish._num(0.949999988079071), "0.95")

    def test_leaderboard_lists_latest_eval_runner_runs(self):
        board = self.files["leaderboard.md"].decode()
        self.assertIn("## Latest eval-runner runs", board)
        self.assertIn("| 2026-10-02 | glm-5.3-flash | GPQA diamond | 50.00% |", board)

    @unittest.skipUnless(openpyxl, "openpyxl not installed")
    def test_xlsx_sheets_rank_and_hold_every_row(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.files["leaderboard.xlsx"]))
        self.assertEqual(wb.sheetnames, ["GPQA (8k)", "IFEval (8k)", "All results", "Notes"])
        gpqa = wb["GPQA (8k)"]
        self.assertEqual(gpqa["A1"].value, "Rank")
        self.assertEqual(gpqa["B2"].value, "glm-5.3-flash")
        self.assertEqual(gpqa["C2"].value, 50.0)
        self.assertEqual(gpqa["C2"].number_format, '0.00"%"')
        self.assertEqual(gpqa.freeze_panes, "C2")
        self.assertIsNotNone(gpqa.auto_filter.ref)
        everything = wb["All results"]
        self.assertEqual(everything.max_row, 1 + len(self.rows))
        self.assertEqual([c.value for c in everything[1]], [t for t, _ in publish.COLUMNS if t != "Key"])


class SeriesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        make_tree(self.tmp.name)
        add_32k_run(self.tmp.name)
        self.files, self.rows = publish.build_files(publish.collect(self.tmp.name), None, "NOW")

    def tearDown(self):
        self.tmp.cleanup()

    def test_32k_row_is_its_own_series(self):
        row = next(r for r in self.rows if "32k" in r["Run"])
        self.assertEqual((row["Series"], row["Comparable"], row["Token budget"]), ("32k", "no", 32768))
        self.assertIn("separate 32k series", row["Why not comparable"])

    def test_leaderboard_ranks_32k_separately(self):
        board = self.files["leaderboard.md"].decode()
        self.assertIn("## GPQA diamond: 32k token budget series, not comparable with 8k", board)
        self.assertNotIn("## IFEval: 32k", board)
        not_comparable = board.split("## Not comparable")[1]
        self.assertNotIn("separate 32k series", not_comparable)

    @unittest.skipUnless(openpyxl, "openpyxl not installed")
    def test_xlsx_has_32k_sheet(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.files["leaderboard.xlsx"]))
        self.assertEqual(wb.sheetnames, ["GPQA (8k)", "GPQA (32k)", "IFEval (8k)", "All results", "Notes"])
        self.assertEqual(wb["GPQA (32k)"].max_row, 2)

    def test_32k_views_filter_on_series(self):
        col_ids = {t: i for i, (t, _) in enumerate(publish.COLUMNS)}
        settings = publish.view_settings(col_ids, "GPQA diamond", {"Series": "32k"})
        self.assertIn({"columnId": col_ids["Series"], "operator": "is-equal", "value": "32k"}, settings["filter"][0])


class PublishTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        make_tree(self.tmp.name)
        self.files, self.rows = publish.build_files(publish.collect(self.tmp.name), None, "NOW")
        self.nc = FakeNextcloud()

    def tearDown(self):
        self.tmp.cleanup()

    def test_first_publish_creates_everything(self):
        table_id, counts = publish.publish(self.nc, self.files, self.rows, share_with="steve")
        self.assertEqual(counts, (5, 0, 0))
        state = self.nc.state
        self.assertEqual([t["title"] for t in state["tables"]], [publish.TABLE_TITLE])
        self.assertEqual([c["title"] for c in state["columns"]], [t for t, _ in publish.COLUMNS])
        self.assertEqual(len(state["rows"]), 5)
        self.assertEqual({v["title"] for v in state["views"]}, {v[0] for v in publish.VIEWS})
        titles = {c["id"]: c["title"] for c in state["columns"]}
        gpqa = next(v for v in state["views"] if v["title"] == "GPQA")
        self.assertEqual(gpqa["filter"], [[{"columnId": next(i for i, t in titles.items() if t == "Task"),
                                             "operator": "is-equal", "value": "GPQA diamond"}]])
        self.assertEqual([titles[c["columnId"]] for c in gpqa["columnSettings"]], publish.VIEW_COLUMNS)
        self.assertEqual(gpqa["sort"], [])  # every column sortable from its header
        bfcl = next(v for v in state["views"] if v["title"] == "BFCL")
        self.assertNotIn("Alt score %", [titles[c["columnId"]] for c in bfcl["columnSettings"]])
        runs = next(v for v in state["views"] if v["title"] == "Recent eval-runner runs")
        self.assertEqual([titles[c["columnId"]] for c in runs["columnSettings"]], publish.RUNS_VIEW_COLUMNS)
        self.assertEqual([(titles[f["columnId"]], f["value"]) for f in runs["filter"][0]], [("Source", "eval-runner")])
        # the table and every view are shared, so the views show up for steve
        self.assertEqual({s["receiver"] for s in state["shares"]}, {"steve"})
        self.assertEqual({s["nodeId"] for s in state["shares"] if s.get("nodeType") == "view"},
                         {v["id"] for v in state["views"]})
        self.assertIn(f"{publish.FOLDER}/leaderboard.md", state["files"])
        self.assertEqual(state["deleted"], [f"{publish.FOLDER}/leaderboard.csv"])
        table = state["tables"][0]
        key_id = next(c["id"] for c in state["columns"] if c["title"] == "Key")
        self.assertEqual(titles[table["columnSettings"][0]["columnId"]], "Model")
        self.assertEqual(table["columnSettings"][-1], {"columnId": key_id, "order": len(publish.COLUMNS) - 1})
        self.assertEqual(sorted(publish.TABLE_ORDER), sorted(t for t, _ in publish.COLUMNS))
        self.assertEqual(table["sort"], [])

    def test_republish_is_idempotent(self):
        publish.publish(self.nc, self.files, self.rows, share_with="steve")
        _, counts = publish.publish(self.nc, self.files, self.rows, share_with="steve")
        self.assertEqual(counts, (0, 0, 5))
        state = self.nc.state
        self.assertEqual((len(state["tables"]), len(state["rows"]), len(state["views"]), len(state["shares"])),
                         (1, 5, len(publish.VIEWS), 1 + len(publish.VIEWS)))

    def test_table_share_has_manage_and_old_read_only_share_is_upgraded(self):
        publish.publish(self.nc, self.files, self.rows, share_with="steve")
        table_share = next(s for s in self.nc.state["shares"] if "tableId" in s)
        self.assertTrue(table_share["permissionManage"])
        table_share["permissionManage"] = False  # as shared before 2026-10-02
        publish.publish(self.nc, self.files, self.rows, share_with="steve")
        self.assertTrue(table_share["permissionManage"])
        self.assertEqual(len(self.nc.state["shares"]), 1 + len(publish.VIEWS))

    def test_old_view_titles_are_renamed_in_place(self):
        table_id = publish.ensure_table(self.nc)
        col_ids = publish.ensure_columns(self.nc, table_id)
        old = self.nc.tables("POST", f"/tables/{table_id}/views", {"title": "Comparable: GPQA", "emoji": "x"})
        publish.ensure_views(self.nc, table_id, col_ids)
        titles = [v["title"] for v in self.nc.state["views"]]
        self.assertEqual(len(titles), len(publish.VIEWS))
        self.assertNotIn("Comparable: GPQA", titles)
        self.assertEqual(next(v for v in self.nc.state["views"] if v["id"] == old["id"])["title"], "GPQA")

    def test_half_configured_view_is_repaired(self):
        table_id, _ = publish.publish(self.nc, self.files, self.rows)
        self.nc.state["views"][0].pop("filter")
        publish.publish(self.nc, self.files, self.rows)
        self.assertIn("filter", self.nc.state["views"][0])
        self.assertEqual(len(self.nc.state["views"]), len(publish.VIEWS))

    def test_mock_rejects_the_old_string_format(self):
        import mock_nextcloud
        self.assertIsNotNone(mock_nextcloud.view_update_problem({"columns": "[1,2]"}))
        self.assertIsNotNone(mock_nextcloud.view_update_problem({"sort": [{"columnId": 1, "mode": "down"}]}))
        self.assertIsNone(mock_nextcloud.view_update_problem(publish.view_settings(
            {t: i for i, (t, _) in enumerate(publish.COLUMNS)}, "IFEval", {"Comparable": "yes"})))

    def test_changed_value_updates_in_place(self):
        publish.publish(self.nc, self.files, self.rows)
        changed = [dict(r) for r in self.rows]
        changed[0]["Note"] = "edited"
        _, counts = publish.publish(self.nc, self.files, changed)
        self.assertEqual(counts, (0, 1, 4))
        self.assertEqual(len(self.nc.state["rows"]), 5)

    def test_numbers_returned_as_floats_are_unchanged(self):
        self.assertTrue(publish._same(198.0, 198))
        self.assertTrue(publish._same("43.94", 43.94))
        self.assertFalse(publish._same(43.0, 43.94))


class ClientTest(unittest.TestCase):
    def test_network_error_becomes_nextcloud_error(self):
        import urllib.error

        def refuse(req, timeout=None):
            raise urllib.error.URLError("connection refused")
        nc = publish.Nextcloud("http://x", "u", "p", opener=refuse)
        with self.assertRaises(publish.NextcloudError):
            nc.tables("GET", "/tables")

    def test_webdav_paths_are_quoted(self):
        nc = publish.Nextcloud("https://nc/", "eval-reports", "p")
        self.assertEqual(nc._dav("Reports/eval-runner/a b.md"),
                         "/remote.php/dav/files/eval-reports/Reports/eval-runner/a%20b.md")


class MainTest(unittest.TestCase):
    def test_dry_run_writes_files(self):
        with tempfile.TemporaryDirectory() as root, tempfile.TemporaryDirectory() as out:
            make_tree(root)
            with unittest.mock.patch("builtins.print"):
                rc = publish.main(["--results-root", root, "--dry-run", out, "--findings", "/nonexistent"])
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.exists(os.path.join(out, "leaderboard.md")))

    def test_unconfigured_publish_exits_2(self):
        with tempfile.TemporaryDirectory() as root:
            with unittest.mock.patch.dict(os.environ, {}, clear=True), \
                    unittest.mock.patch("sys.stderr", new_callable=io.StringIO):
                self.assertEqual(publish.main(["--results-root", root, "--findings", "/nonexistent"]), 2)


if __name__ == "__main__":
    unittest.main()
